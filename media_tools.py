# -*- coding: utf-8 -*-
"""
media_tools.py
أدوات تعديل الصور والفيديو.
"""

from PIL import Image, ImageEnhance
import io
import os
import tempfile


# ============ أدوات الصور ============

def resize_image(image: Image.Image, width: int, height: int) -> Image.Image:
    return image.resize((width, height))


def crop_image(image: Image.Image, left: int, top: int, right: int, bottom: int) -> Image.Image:
    return image.crop((left, top, right, bottom))


def rotate_image(image: Image.Image, angle: float) -> Image.Image:
    return image.rotate(angle, expand=True)


def grayscale_image(image: Image.Image) -> Image.Image:
    return image.convert("L")


def adjust_brightness(image: Image.Image, factor: float) -> Image.Image:
    enhancer = ImageEnhance.Brightness(image)
    return enhancer.enhance(factor)


def flip_image(image: Image.Image, mode: str) -> Image.Image:
    if mode == "horizontal":
        return image.transpose(Image.FLIP_LEFT_RIGHT)
    else:
        return image.transpose(Image.FLIP_TOP_BOTTOM)


def image_to_bytes(image: Image.Image, fmt="PNG") -> bytes:
    buf = io.BytesIO()
    if fmt.upper() == "JPEG" and image.mode in ("RGBA", "P"):
        image = image.convert("RGB")
    image.save(buf, format=fmt)
    return buf.getvalue()


# ============ أدوات الفيديو ============

def _save_uploaded_to_temp(data: bytes, suffix=".mp4") -> str:
    tmp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix)
    tmp.write(data)
    tmp.close()
    return tmp.name


def trim_video(input_path: str, output_path: str, start_sec: float, end_sec: float):
    from moviepy.editor import VideoFileClip
    with VideoFileClip(input_path) as clip:
        sub = clip.subclip(start_sec, end_sec)
        sub.write_videofile(output_path, codec="libx264", audio_codec="aac", logger=None)


def extract_audio(input_path: str, output_path: str):
    from moviepy.editor import VideoFileClip
    with VideoFileClip(input_path) as clip:
        clip.audio.write_audiofile(output_path, logger=None)


def resize_video(input_path: str, output_path: str, width: int, height: int):
    from moviepy.editor import VideoFileClip
    with VideoFileClip(input_path) as clip:
        resized = clip.resize(newsize=(width, height))
        resized.write_videofile(output_path, codec="libx264", audio_codec="aac", logger=None)


# ============ إنشاء ملفات مستندات حقيقية ============

def create_pdf_from_text(text: str, output_path: str, title: str = "مستند"):
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    arabic_font_loaded = False
    for font_path in [r"C:\Windows\Fonts\arial.ttf", r"C:\Windows\Fonts\tahoma.ttf"]:
        if os.path.exists(font_path):
            try:
                pdf.add_font("ArFont", "", font_path, uni=True)
                pdf.set_font("ArFont", size=12)
                arabic_font_loaded = True
                break
            except Exception:
                continue
    if not arabic_font_loaded:
        pdf.set_font("Helvetica", size=12)

    for line in text.split("\n"):
        try:
            pdf.multi_cell(0, 8, line)
        except Exception:
            pdf.multi_cell(0, 8, line.encode("latin-1", "replace").decode("latin-1"))

    pdf.output(output_path)


def create_docx_from_text(text: str, output_path: str):
    from docx import Document
    doc = Document()
    for line in text.split("\n"):
        doc.add_paragraph(line)
    doc.save(output_path)


def create_xlsx_from_text(text: str, output_path: str):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    for i, line in enumerate(text.split("\n"), start=1):
        parts = line.split(",") if "," in line else [line]
        for j, part in enumerate(parts, start=1):
            ws.cell(row=i, column=j, value=part.strip())
    wb.save(output_path)


def create_pptx_from_text(text: str, output_path: str):
    from pptx import Presentation
    prs = Presentation()
    layout = prs.slide_layouts[1]
    chunks = [c.strip() for c in text.split("\n\n") if c.strip()]
    if not chunks:
        chunks = [text.strip() or "شريحة فارغة"]
    for idx, chunk in enumerate(chunks):
        slide = prs.slides.add_slide(layout)
        lines = [l for l in chunk.split("\n") if l.strip()]
        slide.shapes.title.text = lines[0][:60] if lines else f"شريحة {idx + 1}"
        body = slide.placeholders[1]
        body.text_frame.text = "\n".join(lines[1:]) if len(lines) > 1 else chunk
    prs.save(output_path)
